import io
import re
import csv
import json
from pathlib import Path
from typing import List, Tuple, Optional
from datetime import datetime

import pypdf
import docx

from src.schemas.models import DocumentMetadata, DocumentChunk
from src.ingestion.chunker import ContextualChunker


PRESET_SUGGESTED_QUERIES = {
    "cve_2024_8192_kernel_security_advisory": [
        "What is the memory offset and vulnerable function in CVE-2024-8192?",
        "Which upstream Git commit fixed CVE-2024-8192 and which distros are patched?",
        "What sysctl workaround hardens against CVE-2024-8192 without rebooting?",
        "Which Linux kernel versions are affected by the po->rollover race condition?",
    ],
    "q3_enterprise_cloud_infrastructure_financials": [
        "What was total consolidated revenue, operating margin, and AI Cloud growth in Q3?",
        "Compare Q3 revenue and gross margin across Compute, Storage, and AI Cloud segments.",
        "What was Q3 capital expenditure (CAPEX) and which GPU models were procured?",
        "What was the Datacenter Power Usage Effectiveness (PUE) and ending cash liquidity?",
    ],
    "distributed_raft_consensus_architecture_spec": [
        "What is the heartbeat interval and randomized follower election timeout in Raft?",
        "When do nodes trigger log compaction and memory snapshots in Raft?",
        "How does the Raft leader serve linearizable reads via Leader Lease Protocol?",
        "What are the 5 core safety invariants of the Distributed Raft Consensus Engine?",
    ],
    "enterprise_msa_gdpr_compliance_sla": [
        "What is the maximum GDPR breach notification window and encryption standard under Section 3?",
        "What are the standard liability cap vs Data Protection super-cap dollar limits?",
        "What service credit percentages apply when monthly uptime drops below 99.99%?",
        "Which governing law, arbitration forum, and EEA data residency regions apply?",
    ],
    "fda_phase3_oncology_clinical_protocol": [
        "What was the median Progression-Free Survival (PFS) and Hazard Ratio in Cohort A vs Cohort B?",
        "What are the CTCAE v5.0 dosage modification rules for Grade 2 vs Grade 3/4 pneumonitis?",
        "What are the mandatory inclusion criteria and absolute checkpoint inhibitor exclusions?",
        "What was the 12-month Overall Survival (OS) rate and PD-L1 TPS >= 50% subgroup PFS?",
    ],
    "sev1_global_payment_gateway_postmortem": [
        "What caused the 47-minute SEV-1 payment gateway outage and PgBouncer pool exhaustion?",
        "What emergency kubectl, psql, and redis-cli commands mitigated INC-2024-1004?",
        "What were the peak P99 checkout latency and failed transaction counts during the outage?",
        "What are the 3 mandatory P0 remediation action items and assigned SRE team owners?",
    ],
}


class DocumentLoader:
    """Loads documents from directory or raw text strings into structured chunks."""

    def __init__(self, chunker: Optional[ContextualChunker] = None):
        self.chunker = chunker or ContextualChunker()

    @staticmethod
    def extract_text_and_pages(filename: str, raw_bytes: bytes) -> Tuple[str, int, str]:
        """Extracts clean Markdown-structured text, total pages, and doc_type from raw bytes (.pdf, .docx, .csv, .json, .md, .txt)."""
        ext = Path(filename).suffix.lower()

        # 1. PDF
        if ext == ".pdf":
            try:
                reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
                pages_text = []
                total_pages = len(reader.pages)
                for idx, page in enumerate(reader.pages, start=1):
                    p_text = page.extract_text() or ""
                    if p_text.strip():
                        pages_text.append(f"<!-- page: {idx} -->\n## Page {idx}\n\n{p_text.strip()}")
                combined = "\n\n".join(pages_text) if pages_text else f"# {Path(filename).stem}\n\n[Empty or Scanned PDF]"
                return combined, max(total_pages, 1), "pdf"
            except Exception as e:
                return f"# {Path(filename).stem}\n\nError parsing PDF: {e}", 1, "pdf"

        # 2. DOCX
        elif ext in (".docx", ".doc"):
            try:
                doc = docx.Document(io.BytesIO(raw_bytes))
                parts = []
                title_set = False
                for p in doc.paragraphs:
                    t = p.text.strip()
                    if not t:
                        continue
                    style_name = (p.style.name if p.style else "").lower()
                    if "heading 1" in style_name:
                        parts.append(f"# {t}")
                        title_set = True
                    elif "heading 2" in style_name:
                        parts.append(f"## {t}")
                    elif "heading 3" in style_name:
                        parts.append(f"### {t}")
                    elif not title_set and len(parts) == 0:
                        parts.append(f"# {t}")
                        title_set = True
                    else:
                        parts.append(t)

                for table in doc.tables:
                    rows = []
                    col_count = len(table.columns) if table.columns else 0
                    for r in table.rows:
                        row_vals = [c.text.strip().replace("\n", " ").replace("|", "\\|") for c in r.cells]
                        rows.append("| " + " | ".join(row_vals) + " |")
                    if rows and col_count > 0:
                        header_sep = "| " + " | ".join(["---"] * col_count) + " |"
                        table_md = "\n" + rows[0] + "\n" + header_sep + "\n" + "\n".join(rows[1:]) + "\n"
                        parts.append(table_md)

                combined = "\n\n".join(parts) if parts else f"# {Path(filename).stem}\n\n[Empty DOCX Document]"
                return combined, 1, "docx"
            except Exception as e:
                return f"# {Path(filename).stem}\n\nError parsing DOCX: {e}", 1, "docx"

        # 3. CSV
        elif ext == ".csv":
            try:
                decoded = raw_bytes.decode("utf-8", errors="replace")
                reader = csv.reader(io.StringIO(decoded))
                rows = list(reader)
                if rows:
                    header = [c.replace("|", "\\|").replace("\n", " ") for c in rows[0]]
                    col_count = len(header)
                    sep = ["---"] * col_count
                    md_rows = [
                        "| " + " | ".join(header) + " |",
                        "| " + " | ".join(sep) + " |",
                    ]
                    for r in rows[1:]:
                        padded = (r + [""] * col_count)[:col_count]
                        clean_padded = [c.replace("|", "\\|").replace("\n", " ") for c in padded]
                        md_rows.append("| " + " | ".join(clean_padded) + " |")
                    combined = f"# {Path(filename).stem.replace('_', ' ').title()}\n\n" + "\n".join(md_rows)
                else:
                    combined = f"# {Path(filename).stem}\n\nEmpty CSV dataset."
                return combined, 1, "csv"
            except Exception as e:
                return f"# {Path(filename).stem}\n\nError parsing CSV: {e}", 1, "csv"

        # 4. JSON
        elif ext == ".json":
            try:
                decoded = raw_bytes.decode("utf-8", errors="replace")
                parsed = json.loads(decoded)
                pretty = json.dumps(parsed, indent=2)
                title = Path(filename).stem.replace("_", " ").title()
                combined = f"# {title} Data Record\n\n```json\n{pretty}\n```"
                return combined, 1, "json"
            except Exception as e:
                decoded = raw_bytes.decode("utf-8", errors="replace")
                return f"# {Path(filename).stem}\n\n```json\n{decoded}\n```", 1, "json"

        # 5. Default Text / Markdown (.md, .markdown, .txt, or fallback)
        else:
            decoded = raw_bytes.decode("utf-8", errors="replace")
            page_tags = re.findall(r"<!--\s*page:\s*(\d+)\s*-->|\[Page\s*(\d+)\]", decoded, re.IGNORECASE)
            max_page = 1
            for p1, p2 in page_tags:
                val = int(p1 or p2)
                if val > max_page:
                    max_page = val
            doc_type = "markdown" if ext in [".md", ".markdown"] else "text"
            return decoded, max_page, doc_type

    def load_bytes(
        self,
        filename: str,
        content: bytes,
        domain: Optional[str] = None,
        title: Optional[str] = None,
    ) -> Tuple[DocumentMetadata, List[DocumentChunk]]:
        """Parses in-memory bytes of any supported format into metadata and chunks."""
        text, total_pages, doc_type = self.extract_text_and_pages(filename, content)
        doc_stem = Path(filename).stem
        doc_id = re.sub(r"[^a-z0-9]+", "_", doc_stem.lower()).strip("_") or "uploaded_doc"
        doc_title = title or self._extract_title(text, doc_stem)

        chunks = self.chunker.chunk_document(
            document_id=doc_id,
            document_title=doc_title,
            text=text,
            base_metadata={"file_name": filename, "format": doc_type},
        )

        inferred_domain = domain or self._infer_domain(doc_id, text)
        suggested = self.generate_suggested_queries(doc_title, chunks)

        metadata = DocumentMetadata(
            document_id=doc_id,
            title=doc_title,
            file_path=filename,
            document_type=doc_type,
            tags=["user-upload", doc_type, inferred_domain],
            total_pages=total_pages,
            total_chunks=len(chunks),
            created_at=datetime.utcnow().isoformat(),
            domain=inferred_domain,
            suggested_queries=suggested,
        )
        return metadata, chunks

    def load_file(self, file_path: Path) -> Tuple[DocumentMetadata, List[DocumentChunk]]:
        """Reads a file and converts it into metadata and chunks."""
        raw_bytes = file_path.read_bytes()
        metadata, chunks = self.load_bytes(filename=file_path.name, content=raw_bytes)
        metadata.file_path = str(file_path)

        if metadata.document_id in PRESET_SUGGESTED_QUERIES:
            metadata.suggested_queries = PRESET_SUGGESTED_QUERIES[metadata.document_id]

        return metadata, chunks

    def load_directory(self, dir_path: Path) -> List[Tuple[DocumentMetadata, List[DocumentChunk]]]:
        """Scans a directory for all supported enterprise documents in canonical order."""
        results = []
        if not dir_path.exists():
            return results

        preferred_order = [
            "cve_2024_8192_kernel_security_advisory",
            "q3_enterprise_cloud_infrastructure_financials",
            "distributed_raft_consensus_architecture_spec",
            "enterprise_msa_gdpr_compliance_sla",
            "fda_phase3_oncology_clinical_protocol",
            "sev1_global_payment_gateway_postmortem",
        ]

        all_files = [
            p for p in dir_path.glob("*.*")
            if p.suffix.lower() in [".md", ".markdown", ".txt", ".pdf", ".docx", ".csv", ".json"]
        ]

        def sort_key(p: Path):
            stem = p.stem.replace("-", "_").lower()
            if stem in preferred_order:
                return (0, preferred_order.index(stem))
            return (1, stem)

        for path in sorted(all_files, key=sort_key):
            results.append(self.load_file(path))
        return results

    @staticmethod
    def generate_suggested_queries(title: str, chunks: List[DocumentChunk]) -> List[str]:
        """Dynamically generates 4 varied, natural-sounding questions from a document's sections."""
        skip_terms = (
            "overview", "introduction", "walkthrough", "video", "table of contents",
            "license", "acknowledgment", "contributing", "changelog", "summary",
        )
        seen_sections: List[str] = []
        for chunk in chunks:
            raw_sec = chunk.section_title.replace("(Table Data)", "").strip()
            clean_sec = re.sub(r"^(\d+[\.\)]|\d+\s*//)\s*", "", raw_sec).strip()
            if not clean_sec or len(clean_sec) < 4:
                continue
            low = clean_sec.lower()
            if any(term in low for term in skip_terms):
                continue
            if clean_sec not in seen_sections:
                seen_sections.append(clean_sec)

        templates = [
            "How does {sec} work and what are its main rules or numbers?",
            "What are the key details and requirements under {sec}?",
            "Break down the metrics, thresholds, and steps in {sec}.",
        ]
        queries: List[str] = []
        for idx, sec in enumerate(seen_sections[:3]):
            tpl = templates[idx % len(templates)]
            queries.append(tpl.format(sec=sec))

        clean_title = re.sub(r"\s*//.*$", "", title).strip() or title
        queries.append(f"What are the most important numbers, limits, and key takeaways in {clean_title}?")
        return queries[:4]

    def _extract_title(self, text: str, fallback: str) -> str:
        for line in text.split("\n"):
            line = line.strip()
            if line.startswith("# "):
                return line.lstrip("# ").strip()
        return fallback.replace("_", " ").title()

    def _infer_domain(self, doc_id: str, text: str) -> str:
        lower = f"{doc_id} {text[:400]}".lower()
        if "cve" in lower or "kernel" in lower or "vulnerability" in lower:
            return "security"
        if "financial" in lower or "q3" in lower or "revenue" in lower or "capex" in lower:
            return "financials"
        if "raft" in lower or "consensus" in lower or "quorum" in lower:
            return "distributed"
        if "msa" in lower or "gdpr" in lower or "liability" in lower or "indemnification" in lower:
            return "legal"
        if "clinical" in lower or "fda" in lower or "phase iii" in lower or "oncology" in lower:
            return "clinical"
        if "sev-1" in lower or "sev1" in lower or "post-mortem" in lower or "outage" in lower:
            return "sre"
        return "custom"

    def _infer_tags(self, filename: str, text: str) -> List[str]:
        tags = []
        name_lower = f"{filename} {text[:300]}".lower()
        if "cve" in name_lower or "security" in name_lower:
            tags.extend(["security", "cve-advisory", "kernel"])
        if "financial" in name_lower or "q3" in name_lower:
            tags.extend(["financial-10q", "earnings", "capex", "tables"])
        if "raft" in name_lower or "consensus" in name_lower:
            tags.extend(["distributed-rfc", "raft", "consensus"])
        if "msa" in name_lower or "gdpr" in name_lower:
            tags.extend(["legal-msa", "gdpr-dpa", "sla"])
        if "clinical" in name_lower or "fda" in name_lower:
            tags.extend(["fda-phase3", "clinical-trial", "oncology"])
        if "sev1" in name_lower or "outage" in name_lower or "post-mortem" in name_lower:
            tags.extend(["sre-postmortem", "sev-1", "incident-response"])
        return tags or ["enterprise-corpus"]
