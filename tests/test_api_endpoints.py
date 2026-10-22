"""Integration tests for FastAPI endpoints."""

import pytest
from fastapi.testclient import TestClient
from src.api.app import app, app_state


@pytest.fixture(scope="module")
def client():
    # Bootstrap default benchmark documents
    app_state.bootstrap_default_documents()
    return TestClient(app)


def test_health_endpoint(client):
    res = client.get("/api/v1/health")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "healthy"
    assert data["total_documents"] >= 3
    assert data["total_chunks"] >= 5


def test_list_and_get_documents(client):
    res = client.get("/api/v1/documents")
    assert res.status_code == 200
    docs = res.json()
    assert len(docs) >= 3

    first_doc_id = docs[0]["document_id"]
    detail_res = client.get(f"/api/v1/documents/{first_doc_id}")
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert "metadata" in detail
    assert "chunks" in detail
    assert len(detail["chunks"]) > 0


def test_search_endpoint_hybrid_and_cache(client):
    # First search request
    payload = {
        "query": "What is the memory offset and vulnerable function in CVE-2024-8192?",
        "mode": "hybrid",
        "top_k": 5,
        "use_cache": True,
    }
    res1 = client.post("/api/v1/search", json=payload)
    assert res1.status_code == 200
    data1 = res1.json()
    assert "0x7fff90a4" in data1["answer"]
    assert len(data1["citations"]) >= 1
    assert data1["faithfulness_score"] >= 0.8
    assert data1["cache_hit"] is False

    # Second identical query should be CACHE HIT
    res2 = client.post("/api/v1/search", json=payload)
    assert res2.status_code == 200
    data2 = res2.json()
    assert data2["cache_hit"] is True
    assert data2["latency_ms"] < 50.0


def test_search_raw_inspection(client):
    payload = {
        "query": "Q3 consolidated revenue and AI cloud growth",
        "top_k": 5,
    }
    res = client.post("/api/v1/search/raw", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert "timings_ms" in data
    assert "dense_top" in data
    assert "bm25_top" in data
    assert "fused_top" in data
    assert "final_reranked" in data


def test_cache_stats_and_clear(client):
    stats_res = client.get("/api/v1/cache/stats")
    assert stats_res.status_code == 200
    stats = stats_res.json()
    assert "total_queries" in stats

    clear_res = client.post("/api/v1/cache/clear")
    assert clear_res.status_code == 200
    assert "cleared" in clear_res.json()["message"]


def test_benchmarks_endpoint(client):
    res = client.get("/api/v1/benchmarks")
    assert res.status_code == 200
    data = res.json()
    assert "performance_matrix" in data
    matrix = data["performance_matrix"]
    assert "hybrid_reranked" in matrix
    assert matrix["hybrid_reranked"]["avg_context_recall"] > 0.8


def test_auth_detect_endpoint(client):
    res = client.get("/api/v1/auth/detect")
    assert res.status_code == 200
    data = res.json()
    assert "provider" in data
    assert "model" in data
    assert "label" in data
    assert "status" in data


def test_test_key_endpoint(client):
    # Empty key test
    res = client.post("/api/v1/auth/test_key", json={"provider": "gemini", "api_key": ""})
    assert res.status_code == 200
    assert res.json()["success"] is False

    # Unsupported provider
    res2 = client.post("/api/v1/auth/test_key", json={"provider": "unknown_ai", "api_key": "dummy"})
    assert res2.status_code == 200
    assert res2.json()["success"] is False


def test_auth_set_key_status_and_disconnect(client):
    # 1. Set a valid test key
    res = client.post("/api/v1/auth/set_key", json={"provider": "gemini", "model": "gemini-2.0-flash", "api_key": "test_ai_key_1234567890"})
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["provider"] == "gemini"

    # 2. Probe status
    status_res = client.get("/api/v1/auth/key_status")
    assert status_res.status_code == 200
    st = status_res.json()
    assert st["has_key"] is True
    assert st["provider"] == "gemini"
    assert st["status"] == "connected"
    assert "..." in st["masked_key"]

    # 3. Disconnect
    disco_res = client.post("/api/v1/auth/disconnect")
    assert disco_res.status_code == 200
    assert disco_res.json()["success"] is True

    # 4. Probe status after disconnect
    status_after = client.get("/api/v1/auth/key_status")
    assert status_after.status_code == 200
    st_after = status_after.json()
    assert st_after["has_key"] is False
    assert st_after["status"] == "offline"


def test_six_enterprise_domains_and_suggested_queries(client):
    res = client.get("/api/v1/documents")
    assert res.status_code == 200
    docs = res.json()
    assert len(docs) >= 6

    for doc in docs[:6]:
        assert doc.get("domain")
        assert isinstance(doc.get("suggested_queries"), list)
        assert len(doc["suggested_queries"]) >= 3


def test_new_domains_grounded_synthesis(client):
    # Legal MSA / GDPR SLA domain
    r_legal = client.post(
        "/api/v1/search",
        json={
            "query": "What is the service credit percentage if monthly uptime drops to 99.50% and what is the GDPR breach notification window?",
            "mode": "hybrid",
            "use_cache": False,
        },
    )
    assert r_legal.status_code == 200
    d_legal = r_legal.json()
    assert "25%" in d_legal["answer"]
    assert "14 hours" in d_legal["answer"]

    # FDA Phase III Clinical Protocol domain
    r_clin = client.post(
        "/api/v1/search",
        json={
            "query": "What was the median Progression-Free Survival (mPFS) and Hazard Ratio for Cohort A in NCT-09482110?",
            "mode": "hybrid",
            "use_cache": False,
        },
    )
    assert r_clin.status_code == 200
    d_clin = r_clin.json()
    assert "14.8 months" in d_clin["answer"]
    assert "0.54" in d_clin["answer"]

    # Production SRE SEV-1 Post-Mortem domain
    r_sre = client.post(
        "/api/v1/search",
        json={
            "query": "What caused the SEV-1 payment gateway outage INC-2024-1004 and what was the exact rollback command?",
            "mode": "hybrid",
            "use_cache": False,
        },
    )
    assert r_sre.status_code == 200
    d_sre = r_sre.json()
    assert "4,096" in d_sre["answer"]
    assert "kubectl rollout undo" in d_sre["answer"]


def test_custom_document_ingest_and_query_workflow(client):
    ingest_payload = {
        "title": "Zero-Trust KMS Security Specification v5",
        "domain": "Cybersecurity & Kernel Sec",
        "author": "Cloud Security Team",
        "category": "custom_spec",
        "content": (
            "# Zero-Trust KMS Security Specification v5\n\n"
            "## 1. Envelope Encryption & Key Rotation\n"
            "All production envelope keys under `mrk-778899aabb` rotate every **3,600 seconds** "
            "and enforce VPC endpoint `vpce-099118822aabb` with session TTL **900 seconds**.\n"
        ),
    }
    ing_res = client.post("/api/v1/documents/ingest", json=ingest_payload)
    assert ing_res.status_code == 200
    new_doc = ing_res.json()
    assert new_doc["document_id"] == "zero_trust_kms_security_specification_v5"
    assert len(new_doc["suggested_queries"]) >= 3

    # Query the newly ingested custom document
    q_res = client.post(
        "/api/v1/search",
        json={
            "query": "What is the key rotation interval for mrk-778899aabb and VPC endpoint?",
            "mode": "hybrid",
            "document_id": new_doc["document_id"],
            "use_cache": False,
        },
    )
    assert q_res.status_code == 200
    q_data = q_res.json()
    assert "3,600" in q_data["answer"] or "mrk-778899aabb" in q_data["answer"]
    assert len(q_data["citations"]) >= 1


def test_seven_page_threeui_architecture_routes(client):
    expected_pages = [
        ("/", "webgl-vanguard"),
        ("/overview", "webgl-vanguard"),
        ("/search", "bg-waves"),
        ("/pipeline", "canvas-layers"),
        ("/documents", "void-dotmatrix"),
        ("/ingest", "glcanvas-matrix"),
        ("/analytics", "dial-arc-p5"),
        ("/settings", "flow-canvas"),
    ]
    for route, signature in expected_pages:
        res = client.get(route)
        assert res.status_code == 200, f"Route {route} failed with {res.status_code}"
        assert signature in res.text, f"Route {route} missing ThreeUI signature {signature}"
        assert "synapse-top-dock" in res.text


def test_document_deletion_endpoint(client):
    # Ingest a temporary test doc
    payload = {
        "title": "Ephemeral Deletion Test Spec",
        "content": "UniqueEphemeralTokenAlpha9988: Confidential hardware token for deletion verification.",
        "domain": "security",
        "document_type": "markdown",
    }
    ing_res = client.post("/api/v1/documents/ingest", json=payload)
    assert ing_res.status_code == 200
    doc_id = ing_res.json()["document_id"]

    # Verify doc exists
    get_res = client.get(f"/api/v1/documents/{doc_id}")
    assert get_res.status_code == 200

    # Delete doc
    del_res = client.delete(f"/api/v1/documents/{doc_id}")
    assert del_res.status_code == 200
    del_data = del_res.json()
    assert del_data["status"] == "deleted"
    assert del_data["document_id"] == doc_id

    # Verify doc no longer exists
    get_res2 = client.get(f"/api/v1/documents/{doc_id}")
    assert get_res2.status_code == 404


def test_multi_format_upload_endpoint(client):
    import io
    import docx

    # 1. Test CSV Upload
    csv_bytes = b"Server,IP,Region,Status\napi-prod-01,10.0.1.5,us-east-1,ONLINE\napi-prod-02,10.0.1.6,us-east-1,ONLINE"
    res_csv = client.post(
        "/api/v1/documents/upload",
        files={"file": ("servers_inventory.csv", csv_bytes, "text/csv")},
        data={"domain": "sre"},
    )
    assert res_csv.status_code == 200
    doc_csv = res_csv.json()
    assert doc_csv["document_type"] == "csv"
    assert doc_csv["total_chunks"] >= 1

    # 2. Test JSON Upload
    json_bytes = b'{"cluster": "k8s-prod-us", "nodes": 128, "version": "v1.31.1"}'
    res_json = client.post(
        "/api/v1/documents/upload",
        files={"file": ("cluster_config.json", json_bytes, "application/json")},
        data={"domain": "architecture"},
    )
    assert res_json.status_code == 200
    doc_json = res_json.json()
    assert doc_json["document_type"] == "json"
    assert doc_json["total_chunks"] >= 1

    # 3. Test DOCX Upload
    doc = docx.Document()
    doc.add_heading("Production Runbook", level=1)
    doc.add_paragraph("Step 1: Check Prometheus alarms.")
    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    res_docx = client.post(
        "/api/v1/documents/upload",
        files={"file": ("runbook.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"domain": "sre"},
    )
    assert res_docx.status_code == 200
    doc_docx = res_docx.json()
    assert doc_docx["document_type"] == "docx"
    assert doc_docx["total_chunks"] >= 1




